"""Build an evidence-gated AOG quote workbook template; never invent missing prices."""
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
import argparse

def build_quote_template(output_path: str | Path) -> str:
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "AOG Quote"
    ws.append(["NAVARIS — SUCO/GUPCO Ras Shukeir AOG Quote Template"])
    ws.append(["Status", "PRE-QUOTE ONLY — buyer demand and 500 kg scenario not verified"])
    ws.append(["Prepared", "2026-10-16"])
    ws.append(["Route", "Dubai/Jebel Ali → Adabiya/Suez → Ras Shukeir (proposed; verify)"])
    ws.append(["Shipment weight kg", 500])
    ws.append(["Item / part number", "REQUIRED — buyer confirmation"])
    ws.append(["Buyer RFQ / PO", "REQUIRED — buyer confirmation"])
    ws.append(["Supplier/equipment cost USD", None])
    ws.append(["International freight USD", None])
    ws.append(["Customs, duties and taxes USD", None])
    ws.append(["Insurance and handling USD", None])
    ws.append(["Final-mile transport USD", None])
    ws.append(["Other approved direct costs USD", None])
    ws.append(["Agreed service/coordination fee USD", None])
    ws.append(["Total direct costs USD", '=IF(COUNT(B8:B13)<6,"",SUM(B8:B13))'])
    ws.append(["Total proposed quote USD", '=IF(COUNT(B8:B14)<7,"",B15+B14)'])
    ws.append(["Licensed operator/CITL proposed share (90% of fee)", '=IF(ISNUMBER(B14),B14*90%,"")'])
    ws.append(["NAVARIS proposed share (10% of fee)", '=IF(ISNUMBER(B14),B14*10%,"")'])
    ws.append(["Supplier stock / quote reference", "REQUIRED"])
    ws.append(["Freight quote / validity", "REQUIRED"])
    ws.append(["Delivery window confirmed in writing", "REQUIRED — do not promise 72h before confirmation"])
    ws.append(["Payment terms / trigger", "REQUIRED — signed written agreement"])
    ws.append(["Evidence status", "HOLD — no buyer RFQ, supplier quote or logistics quote evidenced"])
    ws.append(["Commercial note", "90/10 is a proposal on the agreed service fee only, not equipment value; sign a written schedule first."])
    ws.freeze_panes = "A8"
    ws.column_dimensions["A"].width = 48
    ws.column_dimensions["B"].width = 82
    ws.auto_filter.ref = "A1:B24"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].fill = PatternFill("solid", fgColor="D9EAF7")
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if cell.column == 1:
                cell.font = Font(bold=True)
    for r in (15,16,17,18):
        ws.cell(r,1).fill = PatternFill("solid", fgColor="FFF2CC")
        ws.cell(r,2).fill = PatternFill("solid", fgColor="FFF2CC")
    dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
    dv.error = "Enter a non-negative numeric amount or leave blank until verified."
    dv.errorTitle = "Invalid cost input"
    ws.add_data_validation(dv)
    dv.add("B8:B14")
    ws.sheet_view.showGridLines = False
    wb.save(path)
    # Reopen to verify the file is valid and the pricing formulas are present.
    check = load_workbook(path, data_only=False)
    assert "AOG Quote" in check.sheetnames
    assert str(check["AOG Quote"]["B16"].value).startswith("=IF(")
    return str(path)

def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", default="output/navaris-quote-suco-aog-500kg.xlsx")
    args = parser.parse_args()
    result = build_quote_template(args.output)
    print(f"NAVARIS quote template created: {result}")
    print("No prices are populated; do not issue as a commercial quotation until inputs are verified.")
    return 0

if __name__ == "__main__":
    raise SystemExit(main())
