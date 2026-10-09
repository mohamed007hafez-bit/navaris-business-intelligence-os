from pathlib import Path
from openpyxl import load_workbook

WORKBOOK = Path("data/NAVARIS_3Part_AgreementFirst_OS.xlsx")
MASTER = "MASTER_3Part_Agreement_First"
OPPORTUNITIES = "OPPORTUNITY_MAP"

def fail(message):
    raise SystemExit(f"FAIL: {message}")

if not WORKBOOK.exists():
    fail(f"workbook not found: {WORKBOOK}")

wb = load_workbook(WORKBOOK, data_only=False)
if MASTER not in wb.sheetnames:
    fail(f"missing required sheet: {MASTER}")
if OPPORTUNITIES not in wb.sheetnames:
    fail(f"missing required sheet: {OPPORTUNITIES}")

ws = wb[MASTER]
headers = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
if not any(headers):
    fail("master sheet header row is empty")
if "Agreement_Status" not in headers:
    fail("missing Agreement_Status control column")

master_rows = [
    [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
    for r in range(2, ws.max_row + 1)
]
active_master_rows = [row for row in master_rows if any(v not in (None, "") for v in row)]
if not active_master_rows:
    fail("master sheet contains no active opportunity rows")

op = wb[OPPORTUNITIES]
opportunity_count = sum(
    1 for row in op.iter_rows(min_row=2, values_only=True)
    if any(v not in (None, "") for v in row)
)
if opportunity_count < 10:
    fail(f"expected at least 10 opportunity-map rows; found {opportunity_count}")

print("PASS: workbook exists and opens.")
print(f"PASS: required sheets present: {MASTER}, {OPPORTUNITIES}.")
print(f"PASS: Agreement_Status control column present; active master rows={len(active_master_rows)}.")
print(f"PASS: opportunity-map rows={opportunity_count} (minimum 10).")
print("NOTICE: this is a structural test, not proof that demand, suppliers, budgets, or source evidence are verified.")
print("NOTICE: CASH COLLECTED must remain zero unless actual receipt evidence is recorded.")
