"""Financial and Excel analysis helpers using pandas/openpyxl."""
from __future__ import annotations
from pathlib import Path
from typing import Any
import pandas as pd
from openpyxl import Workbook, load_workbook
from navaris_agents.commercial_controls import verify_bank_receipt


def summarize_cashflow(rows: list[dict[str, Any]]) -> dict[str, Any]:
    """A boolean flag alone never counts as cash; full BANK_RECEIPT audit metadata is required."""
    if not rows:
        return {"status":"NO_DATA","cash_in":0.0,"cash_out":0.0,"net_cashflow":0.0,"cash_collected":0.0,"row_count":0}
    frame=pd.DataFrame(rows)
    required={"amount","direction"}
    missing=required-set(frame.columns)
    if missing: raise ValueError(f"Missing required columns: {', '.join(sorted(missing))}")
    frame["amount"]=pd.to_numeric(frame["amount"],errors="raise")
    if (frame["amount"]<0).any(): raise ValueError("amount must be non-negative; direction defines inflow/outflow")
    direction=frame["direction"].astype(str).str.lower()
    if (~direction.isin({"in","out"})).any(): raise ValueError("direction must be 'in' or 'out'")
    cash_in=float(frame.loc[direction=="in","amount"].sum())
    cash_out=float(frame.loc[direction=="out","amount"].sum())
    collected=0.0
    for row in rows:
        if str(row.get("direction","")).lower()=="in":
            verified=verify_bank_receipt(row)
            if verified["status"]=="RECEIVED_VERIFIED": collected+=verified["cash_collected"]
    return {"status":"CALCULATED_FROM_INPUT","cash_in":round(cash_in,2),"cash_out":round(cash_out,2),
            "net_cashflow":round(cash_in-cash_out,2),"cash_collected":round(collected,2),"row_count":len(rows),
            "warning":"Collected cash requires BANK_RECEIPT evidence, redacted reference, amount, currency, received date, verifier and verification timestamp."}


def profitability(*, revenue: float, direct_costs: float, operating_costs: float = 0.0) -> dict[str, float]:
    values={"revenue":revenue,"direct_costs":direct_costs,"operating_costs":operating_costs}
    if any(not isinstance(v,(int,float)) or v<0 for v in values.values()):
        raise ValueError("revenue and costs must be non-negative numbers")
    gross_profit=revenue-direct_costs
    operating_profit=gross_profit-operating_costs
    return {"revenue":round(float(revenue),2),"direct_costs":round(float(direct_costs),2),
            "gross_profit":round(float(gross_profit),2),"gross_margin_pct":round(gross_profit/revenue*100,2) if revenue else 0.0,
            "operating_costs":round(float(operating_costs),2),"operating_profit":round(float(operating_profit),2),
            "operating_margin_pct":round(operating_profit/revenue*100,2) if revenue else 0.0}


def read_excel_records(path: str | Path, sheet_name: str | int = 0) -> list[dict[str, Any]]:
    file_path=Path(path)
    if not file_path.is_file() or file_path.suffix.lower() not in {".xlsx",".xlsm"}:
        raise ValueError("Provide an existing .xlsx or .xlsm workbook path")
    frame=pd.read_excel(file_path,sheet_name=sheet_name,engine="openpyxl")
    return frame.where(pd.notna(frame),None).to_dict(orient="records")


def write_analysis_workbook(output: str | Path, *, cash_rows: list[dict[str, Any]], profit_inputs: dict[str, float]) -> str:
    target=Path(output); target.parent.mkdir(parents=True,exist_ok=True)
    cash=summarize_cashflow(cash_rows); profit=profitability(**profit_inputs)
    wb=Workbook(); ws=wb.active; ws.title="Cashflow"; ws.append(["Metric","Value"])
    for key,value in cash.items():
        if isinstance(value,(str,int,float)): ws.append([key,value])
    ps=wb.create_sheet("Profitability"); ps.append(["Metric","Value"])
    for key,value in profit.items():
        if isinstance(value,(str,int,float)): ps.append([key,value])
    for sheet in wb.worksheets:
        sheet.freeze_panes="A2"; sheet.column_dimensions["A"].width=32; sheet.column_dimensions["B"].width=48
    wb.save(target)
    check=load_workbook(target,read_only=True)
    assert {"Cashflow","Profitability"}.issubset(set(check.sheetnames))
    check.close()
    return str(target)
