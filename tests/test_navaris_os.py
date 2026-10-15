import os
import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

import navaris_core as core
from cash_agent import calc_monthly_cash
from chatgpt_bridge import should_call_cloud
from ncnda_agent import generate_ncnda
from pricer_agent import price_route
from scout_agent import scout_tenders


class NavarisCoreTests(unittest.TestCase):
    def test_cost_formula(self):
        self.assertAlmostEqual(core.calc_cost(45, fuel_rate=15, toll=200, surcharge=0), 875)

    def test_profit_and_split_reconcile(self):
        profit = core.calc_profit(1500, 875)
        navaris_share, company_share, company_net = core.calc_share_30_70(1500, 875)
        self.assertAlmostEqual(profit, 625)
        self.assertAlmostEqual(navaris_share + company_net, profit)
        self.assertAlmostEqual(company_share, 1050)

    def test_margin_grading_thresholds(self):
        self.assertEqual(core.evaluate_margin(19.99), "غير مقبول")
        self.assertEqual(core.evaluate_margin(20), "مقبول")
        self.assertEqual(core.evaluate_margin(30), "جيد")
        self.assertEqual(core.evaluate_margin(40), "جيد جدا")
        self.assertEqual(core.evaluate_margin(50), "ممتاز")

    def test_negative_inputs_rejected(self):
        with self.assertRaises(ValueError):
            core.calc_cost(-1)

    def test_scout_never_calls_example_verified_by_wording(self):
        result = scout_tenders()
        self.assertEqual(result.iloc[0]["Decision"], "VERIFY")
        self.assertIn("ILLUSTRATIVE", result.iloc[0]["Evidence Status"])

    def test_scout_requires_explicit_verified_flag(self):
        result = scout_tenders([{
            "name": "Test opportunity", "value": 1000, "distance": 1,
            "selling": 100, "evidence_verified": True, "source_status": "SOURCE CHECKED"
        }])
        self.assertEqual(result.iloc[0]["Decision"], "GO")

    def test_cash_projection_is_not_cash_collected(self):
        df, projection = calc_monthly_cash()
        self.assertAlmostEqual(projection["5 clients (illustrative base)"], 27900)
        self.assertEqual(projection["Actual cash collected (evidence-based)"], 0)
        self.assertIn("Actual Cash Collected", df.columns)

    def test_ncnda_is_draft_only(self):
        draft = generate_ncnda("Counterparty Ltd", "ممتاز")
        self.assertIn("DRAFT FOR LEGAL REVIEW", draft)
        self.assertIn("NOT EXECUTABLE", draft)
        self.assertIn("30.00%", draft)

    def test_cloud_requires_grade_manual_approval_and_opt_in(self):
        old = os.environ.get("NAVARIS_ALLOW_CLOUD")
        try:
            os.environ.pop("NAVARIS_ALLOW_CLOUD", None)
            self.assertFalse(should_call_cloud("ممتاز", manual_approval=True))
            os.environ["NAVARIS_ALLOW_CLOUD"] = "1"
            self.assertFalse(should_call_cloud("جيد جدا", manual_approval=True))
            self.assertFalse(should_call_cloud("ممتاز", manual_approval=False))
            self.assertTrue(should_call_cloud("ممتاز", manual_approval=True))
        finally:
            if old is None:
                os.environ.pop("NAVARIS_ALLOW_CLOUD", None)
            else:
                os.environ["NAVARIS_ALLOW_CLOUD"] = old

    def test_excel_workbook_has_real_charts(self):
        pricing = pd.DataFrame([price_route("Test route", 1, 0, 0, 100)])
        monthly, _ = calc_monthly_cash()
        with tempfile.TemporaryDirectory() as temp:
            target = Path(temp) / "navaris-test.xlsx"
            core.build_excel_with_charts(pricing, monthly, target)
            wb = load_workbook(target)
            self.assertIn("Pricing_Grading", wb.sheetnames)
            self.assertIn("Monthly_30pct", wb.sheetnames)
            self.assertIn("Tenders_BCG", wb.sheetnames)
            self.assertGreaterEqual(len(wb["Pricing_Grading"]._charts), 1)
            self.assertGreaterEqual(len(wb["Monthly_30pct"]._charts), 1)


if __name__ == "__main__":
    unittest.main(verbosity=2)
