import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from navaris_agents import financial, oilfield, shipping
from navaris_agents.router import route_request


class ShippingAgentTests(unittest.TestCase):
    def test_landed_cost_reconciles(self):
        result = shipping.calculate_landed_cost(shipping.ShippingCost(
            goods_value=1000, freight=100, insurance=10, origin_charges=20,
            destination_charges=30, duties_and_taxes=40, other_costs=5), margin_pct=10)
        self.assertEqual(result["landed_cost"], 1205)
        self.assertEqual(result["indicative_sell_price"], 1325.5)
        self.assertEqual(result["status"], "ESTIMATE_ONLY")

    def test_negative_shipping_cost_rejected(self):
        with self.assertRaises(ValueError):
            shipping.calculate_landed_cost(shipping.ShippingCost(goods_value=-1, freight=0))

    def test_document_checklist_mode(self):
        self.assertIn("Bill of lading instructions", shipping.documentation_checklist("sea")["items"])
        with self.assertRaises(ValueError):
            shipping.documentation_checklist("space")


class OilfieldAgentTests(unittest.TestCase):
    def test_tender_defaults_to_hold_without_evidence(self):
        result = oilfield.assess_tender(title="Test", source_url="https://example.com/tender",
            deadline=None, scope_known=False, eligibility_known=False, budget_known=False,
            documents_available=False)
        self.assertEqual(result["decision"], "HOLD_VERIFY")

    def test_expired_tender(self):
        result = oilfield.assess_tender(title="Expired", source_url="https://example.com/tender",
            deadline="2000-01-01", scope_known=True, eligibility_known=True, budget_known=True,
            documents_available=True, source_checked=True)
        self.assertEqual(result["decision"], "EXPIRED")

    def test_risk_register_score(self):
        risks = oilfield.project_risk_register([{"description": "Late materials", "likelihood": 4, "impact": 5}])
        self.assertEqual(risks[0]["score"], 20)
        self.assertEqual(risks[0]["rating"], "HIGH")


class FinancialAgentTests(unittest.TestCase):
    def test_cash_collected_requires_full_bank_receipt_evidence(self):
        result = financial.summarize_cashflow([
            {"amount": 100, "direction": "in", "receipt_verified": True},
            {"amount": 50, "direction": "in", "receipt_verified": False},
            {"amount": 40, "direction": "out"},
            {"amount": 125, "direction": "in", "evidence_type": "BANK_RECEIPT",
             "evidence_reference_redacted": "RCPT-REDACTED-001", "opportunity_id": "NAV-OPP-001",
             "currency": "EGP", "received_date": "2026-10-27T10:00:00+03:00",
             "verified_by": "reviewer-id", "verified_at": "2026-10-27T10:05:00+03:00"},
        ])
        self.assertEqual(result["cash_in"], 275)
        self.assertEqual(result["cash_out"], 40)
        self.assertEqual(result["net_cashflow"], 235)
        self.assertEqual(result["cash_collected"], 125)

    def test_profitability(self):
        result = financial.profitability(revenue=1000, direct_costs=600, operating_costs=100)
        self.assertEqual(result["gross_profit"], 400)
        self.assertEqual(result["operating_profit"], 300)
        self.assertEqual(result["gross_margin_pct"], 40)

    def test_excel_output_is_valid(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "analysis.xlsx"
            output = financial.write_analysis_workbook(path, cash_rows=[
                {"amount": 200, "direction": "in", "receipt_verified": True}],
                profit_inputs={"revenue": 500, "direct_costs": 250, "operating_costs": 50})
            self.assertTrue(Path(output).is_file())
            wb = load_workbook(output, read_only=True)
            self.assertEqual(set(wb.sheetnames), {"Cashflow", "Profitability"})
            wb.close()


class RouterTests(unittest.TestCase):
    def test_disabled_agents_fail_closed(self):
        with self.assertRaises(PermissionError):
            route_request("shipping", "documents", {"mode": "sea"})

    def test_enabled_local_agent_routes_without_external_calls(self):
        result = route_request("shipping", "documents", {"mode": "sea"},
            enabled_agents={"shipping": True}, allow_cloud_ai=False)
        self.assertEqual(result["external_calls"], 0)
        self.assertEqual(result["execution_mode"], "LOCAL_ONLY")

    def test_cloud_requires_human_approval_and_has_no_implicit_provider(self):
        with self.assertRaises(PermissionError):
            route_request("financial", "profitability", {"revenue": 1, "direct_costs": 0},
                enabled_agents={"financial": True}, allow_cloud_ai=True)
        with self.assertRaises(NotImplementedError):
            route_request("financial", "profitability", {"revenue": 1, "direct_costs": 0},
                enabled_agents={"financial": True}, allow_cloud_ai=True, human_approved=True)


if __name__ == "__main__":
    unittest.main(verbosity=2)
